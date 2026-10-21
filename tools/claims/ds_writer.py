"""
ds_writer.py  -  Draws the "Distribution Wise Discount Amount" workbook (layout copied from your sample):
   sheet "Credit Note" = title, heading row, one row per distributor, Total row
   sheet "Check"       = everything to verify before using the summary
Totals are real Excel formulas (=SUM per row, SUBTOTAL(9, ...) in the Total row).
"""
import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from tools.claims import config
from tools.claims import processing
from tools.claims import report_writer

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREY = PatternFill("solid", fgColor="D9D9D9")
YELLOW = PatternFill("solid", fgColor="FFFF99")
ORANGE = PatternFill("solid", fgColor="FFC000")
ACCOUNTING = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'
MONTH_FMT = '[$-409]mmmm\\-yyyy;@'
WIDTHS = [5.1, 10.0, 28.6, 14.7, 14.7, 14.7, 14.7, 14.7, 16.1, 14.7, 14.7, 14.7, 18.9, 16.1]   # A..N (from the sample)


def _f(size):
    return Font(name=config.FONT_NAME, size=size, bold=True)


def _cell(ws, ref, value=None, size=12, h="center", wrap=False, fill=None, fmt=None):
    c = ws[ref]
    indent = 1 if h == "left" else 0
    if value is not None:
        c.value = value
    c.font = _f(size)
    c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap, indent=indent)
    c.border = BOX
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def _summary_sheet(ws, rows, unit, year, month):
    for i, w in enumerate(WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w
    # ---- title + headings ----
    for i in range(14):
        _cell(ws, f"{chr(65 + i)}1", fill=GREY, size=14, wrap=True)
    ws["A1"].value = config.DS_TITLE
    ws.merge_cells("A1:N1")
    for i, text in enumerate(config.DS_HEADINGS):
        _cell(ws, f"{chr(65 + i)}2", text, size=10, wrap=True, fill=GREY)
    ws.row_dimensions[1].height = 35.1
    ws.row_dimensions[2].height = 50.1

    first = 3
    last = first + max(len(rows), 1) - 1
    month_cell = f"$N${first}"
    for k, r in enumerate(rows):
        n = first + k
        ws.row_dimensions[n].height = 35.1
        _cell(ws, f"A{n}", k + 1)
        _cell(ws, f"B{n}", int(r.code) if r.code.isdigit() else r.code)
        _cell(ws, f"C{n}", r.name, h="left", wrap=True)
        for col, val in zip("DEFGHIJK", (r.card, None, r.hth, None, None, r.freight, r.wholesale, None)):
            _cell(ws, f"{col}{n}", val, h=None, fmt=ACCOUNTING)           # empty cells stay empty (fill by hand)
        _cell(ws, f"L{n}", f"=SUM(D{n}:K{n})", h=None, fmt=ACCOUNTING)
        _cell(ws, f"M{n}", unit, h="left")
        _cell(ws, f"N{n}", datetime.date(year, month, 1) if k == 0 else f"={month_cell}", fmt=MONTH_FMT)
    if not rows:                                                            # keep the sheet valid if nothing was found
        for col in "ABCDEFGHIJKLMN":
            _cell(ws, f"{col}{first}")
    # ---- Total row ----
    t = last + 1
    ws.row_dimensions[t].height = 35.1
    for col in "ABC":
        _cell(ws, f"{col}{t}", fill=GREY)
    ws[f"A{t}"].value = "Total"
    ws.merge_cells(f"A{t}:C{t}")
    for col in "DEFGHIJKL":
        _cell(ws, f"{col}{t}", f"=SUBTOTAL(9,{col}{first}:{col}{last})", h=None, fill=GREY, fmt=ACCOUNTING)
    _cell(ws, f"M{t}", unit, h="left")
    _cell(ws, f"N{t}", f"={month_cell}", fmt=MONTH_FMT)
    report_writer._page_setup(ws, t, "N", "1:2", "landscape",
                              f"{unit} {config.DS_FILE_WORD} - {processing.date_label(year, month)}")


def _check_sheet(ws, issues, run_notes, sources, year, month):
    for col, w in zip("ABCD", (30, 12, 34, 110)):
        ws.column_dimensions[col].width = w
    ws["A1"].value = f"Check sheet - {config.DS_FILE_WORD} {processing.date_label(year, month)}"
    ws["A1"].font = Font(name=config.FONT_NAME, size=14, bold=True)
    ws["A2"].value = f"Created: {datetime.datetime.now():%d-%b-%Y %H:%M}"
    ws["A2"].font = Font(name=config.FONT_NAME, size=9)
    ws["A3"].value = ("Rules: Card Discount and HTH = the printed 'Total:' of each distributor's PDF page. "
                      "Freight and Wholesale = Cases x Rate from the reports made by this tool. Distributors are matched by code. "
                      "Empty cells = no figure found (or Non GTM / Periodic / Promo / Special Approval: fill by hand).")
    ws["A3"].font = Font(name=config.FONT_NAME, size=9)

    def line(r, vals, fill=None):
        for col, v in zip("ABCD", vals):
            c = ws[f"{col}{r}"]
            c.value = v
            c.font = Font(name=config.FONT_NAME, size=10, bold=fill is None or fill is GREY)
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if fill:
                c.fill = fill

    line(5, ("Distributor", "Code", "Type", "Details"), GREY)
    r = 6
    for kind, text in run_notes:
        line(r, ("(whole run)", "", kind, text), ORANGE if kind.startswith("WARNING") else YELLOW)
        r += 1
    for name, code, kind, text in issues:
        line(r, (name, code, kind, text), ORANGE if kind.startswith("WARNING") else YELLOW)
        r += 1
    if r == 6:
        ws["A6"].value = "No issues found."
        ws["A6"].font = Font(name=config.FONT_NAME, size=10, bold=True)
        r = 7
    r += 1
    ws[f"A{r}"].value = "Files read"
    ws[f"A{r}"].font = Font(name=config.FONT_NAME, size=11, bold=True)
    r += 1
    for s in sources:
        line(r, ("", "", "Source", s))
        r += 1
    report_writer._page_setup(ws, r, "D", "5:5", "landscape", "Check sheet")


def write_summary_workbook(rows, issues, run_notes, sources, unit, year, month, path):
    wb = Workbook()
    ws = wb.active
    ws.title = config.DS_SHEET_NAME
    _summary_sheet(ws, rows, unit, year, month)
    _check_sheet(wb.create_sheet("Check"), issues, run_notes, sources, year, month)
    wb.calculation.fullCalcOnLoad = True            # Excel recalculates formulas when opened
    wb.save(path)
