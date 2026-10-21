"""
report_writer.py  -  Draws the printable Excel sheets (layout copied from the sample file)
and the "Check" sheet. Totals and amounts are real Excel formulas.
"""
import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

from tools.claims import config
from tools.claims import processing

THIN = Side(style="thin")
MEDIUM = Side(style="medium")
GREY = PatternFill("solid", fgColor="C0C0C0")
YELLOW = PatternFill("solid", fgColor="FFFF99")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _font(size=10, bold=True):
    return Font(name=config.FONT_NAME, size=size, bold=bold)


def _put(ws, ref, value=None, *, fill=None, align=CENTER, fmt=None, size=10, box=True):
    """Write one cell with the standard look."""
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = _font(size)
    c.alignment = align
    if box:
        c.border = BOX
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def _page_setup(ws, last_row, last_col, title_rows, orientation, footer_left):
    """A4, fit to one page wide, print area, margins, centred, repeated header, page numbers."""
    ws.page_setup.paperSize = config.PAPER_SIZE_A4
    ws.page_setup.orientation = orientation
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0                       # rows may flow onto more pages
    ws.print_area = f"A1:{last_col}{last_row}"
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = config.MARGIN_LEFT_RIGHT
    ws.page_margins.top = ws.page_margins.bottom = config.MARGIN_TOP_BOTTOM
    ws.page_margins.header = ws.page_margins.footer = 0.3
    ws.print_title_rows = title_rows
    ws.oddFooter.left.text = footer_left
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddFooter.right.text = "Printed: &D"
    for part in (ws.oddFooter.left, ws.oddFooter.center, ws.oddFooter.right):
        part.size = 8
        part.font = config.FONT_NAME
    ws.sheet_view.showGridLines = False


def _distributor_sheet(ws, rd, year, month):
    n = len(rd.rows)
    first, last = 6, 5 + n          # pack rows
    total = last + 1                # total row
    for col, w in zip("ABCDEF", (13.3, 26.7, 16.1, 13.4, 9.6, 9.7)):
        ws.column_dimensions[col].width = w

    # ---- title ----
    _put(ws, "A1", config.REPORT_TITLE, fill=GREY, size=16)
    for col in "BCDEF":
        _put(ws, f"{col}1", fill=GREY)
    ws.merge_cells("A1:F1")

    # ---- header block (rows 2-4) ----
    left = {2: ("Unit", rd.unit), 3: ("Route", rd.route)}
    for r, (label, val) in left.items():
        _put(ws, f"A{r}", label, fill=GREY, align=Alignment(vertical="center"))
        _put(ws, f"B{r}", val); _put(ws, f"C{r}"); ws.merge_cells(f"B{r}:C{r}")
    _put(ws, "A4", "Month & Year", fill=GREY, align=Alignment(vertical="center"))
    _put(ws, "B4", datetime.date(year, month, 1), fmt="mmm-yy"); _put(ws, "C4"); ws.merge_cells("B4:C4")

    right = {2: ("Area Incharge", rd.area_incharge), 3: ("Distributor", rd.name)}
    for r, (label, val) in right.items():
        _put(ws, f"D{r}", label, fill=GREY, align=Alignment(vertical="center"))
        _put(ws, f"E{r}", val, align=CENTER_WRAP); _put(ws, f"F{r}"); ws.merge_cells(f"E{r}:F{r}")
    _put(ws, "D4", "Total Amount", fill=GREY, align=Alignment(vertical="center"))
    _put(ws, "E4", f"=F{total}", fmt="0"); _put(ws, "F4"); ws.merge_cells("E4:F4")

    # ---- table heading (row 5) ----
    for col, text in zip("ABCDEF", ("Distributor Code", "Description", "Pack", "Rate", "Cases", "Amount")):
        _put(ws, f"{col}5", text, fill=GREY, align=CENTER_WRAP)

    # ---- pack rows ----
    desc = processing.description_text(rd.rate, year, month)
    _put(ws, f"A{first}", int(rd.code) if rd.code.isdigit() else rd.code)
    _put(ws, f"B{first}", desc, align=CENTER_WRAP)
    for r in range(first + 1, last + 1):             # borders for the merged blocks
        _put(ws, f"A{r}"); _put(ws, f"B{r}")
    for i, (label, cases) in enumerate(rd.rows):
        r = first + i
        _put(ws, f"C{r}", label, align=LEFT_WRAP)
        _put(ws, f"D{r}", rd.rate)
        _put(ws, f"E{r}", cases, fmt="#,##0")
        _put(ws, f"F{r}", f"=E{r}*D{r}", fmt="0")
    if n > 1:
        ws.merge_cells(f"A{first}:A{last}")
        ws.merge_cells(f"B{first}:B{last}")

    # ---- total row ----
    _put(ws, f"A{total}", "Total", align=CENTER)
    for col in "BCD":
        _put(ws, f"{col}{total}")
    ws.merge_cells(f"A{total}:D{total}")
    _put(ws, f"E{total}", f"=SUM(E{first}:E{last})", fmt="#,##0")
    _put(ws, f"F{total}", f"=SUM(F{first}:F{last})", fmt="0")

    for r in range(1, total + 1):
        ws.row_dimensions[r].height = 24.95

    # ---- signature blocks ----
    def signatures(line_row, labels):
        ws.row_dimensions[line_row].height = 15.75
        for col, text in zip(("A", "C", "E"), labels):
            targets = [col] if col != "E" else ["E", "F"]
            for t in targets:
                ws[f"{t}{line_row}"].border = Border(bottom=MEDIUM)
            _put(ws, f"{col}{line_row + 1}", text, align=CENTER, box=False)
        ws.merge_cells(f"E{line_row + 1}:F{line_row + 1}")

    signatures(total + 4, config.SIGNATURES_ROW_1)
    signatures(total + 11, config.SIGNATURES_ROW_2)

    _page_setup(ws, total + 12, "F", "5:5", config.ORIENTATION,
                f"{config.REPORT_TITLE} - {processing.date_label(year, month)}")


def _check_sheet(ws, reports, load_notes, year, month, source_name):
    """One line per thing the user should look at before printing."""
    for col, w in zip("ABCDE", (28, 12, 22, 90, 4)):
        ws.column_dimensions[col].width = w
    ws["A1"].value = f"Check sheet - Freight Charges {processing.date_label(year, month)}"
    ws["A1"].font = _font(14)
    ws["A2"].value = f"Source file: {source_name}      Created: {datetime.datetime.now():%d-%b-%Y %H:%M}"
    ws["A2"].font = _font(9, bold=False)
    for col, text in zip("ABCD", ("Distributor", "Code", "Type", "Details")):
        _put(ws, f"{col}4", text, fill=GREY)
    r = 5
    for rep in reports:
        for kind, text in rep.issues:
            for col, val in zip("ABCD", (rep.name, rep.code, kind, text)):
                _put(ws, f"{col}{r}", val, align=LEFT_WRAP, fill=YELLOW)
            r += 1
    for text in load_notes:
        for col, val in zip("ABCD", ("(whole file)", "", "File note", text)):
            _put(ws, f"{col}{r}", val, align=LEFT_WRAP, fill=YELLOW)
        r += 1
    if r == 5:
        _put(ws, "A5", "No issues found.", align=Alignment(horizontal="left"), box=False)
        r = 6
    _page_setup(ws, r, "D", "4:4", "landscape", "Check sheet")


def write_workbook(reports, year, month, path, source_name="", load_notes=()):
    wb = Workbook()
    wb.remove(wb.active)
    for rd in reports:
        ws = wb.create_sheet(rd.sheet_name)
        _distributor_sheet(ws, rd, year, month)
    _check_sheet(wb.create_sheet("Check"), reports, list(load_notes), year, month, source_name)
    wb.calculation.fullCalcOnLoad = True            # Excel recalculates formulas when opened
    wb.save(path)
