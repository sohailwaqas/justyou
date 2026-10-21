"""
wcd_writer.py  -  Draws the "Distribution Wise Monthly Wholesale Claims Detail" workbook (layout copied from your sample):
   sheet "Dic. Cs" = title, Claim Month / Unit / Region, one 7-column block per pack, one row per distributor, TOTAL row
   sheet "Check"   = everything to verify before using the report
Percentages, Total Discounted Cases and totals are real Excel formulas.
Colours: BLUE = Committed Discount that is the sum of the Card Discount and HTH figures (written as =card+hth);
         RED  = "% age of W.S Claim" above the limit (conditional format, so it follows the numbers if you edit them).
"""
import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from tools.claims import config
from tools.claims import processing
from tools.claims import report_writer
from tools.claims.ds_writer import GREY, ORANGE, YELLOW

MED = Side(style="medium")
BOX = Border(left=MED, right=MED, top=MED, bottom=MED)
YEL = PatternFill("solid", fgColor="FFFF00")
BLOCK_FILLS = ["FFFF00", "D6DCE5", "C5E0B4"]                 # header colours of pack block 1, 2, 3 (as in the sample)
TITLE_FILL = PatternFill("solid", fgColor="0D0D0D")
ACC = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'
ACC2 = '_-* #,##0_-;\\-* #,##0_-;_-* "-"_-;_-@_-'
PCT = "0.0%"
WIDTHS = [11.6, 22.7, 15.3, 11.7, 17.4, 11.7, 15.0, 21.0, 11.7, 15.6, 11.7, 16.1, 11.7, 14.3, 17.4,
          11.7, 13.6, 11.7, 13.6, 11.7, 14.4, 15.3, 11.7]     # A..W as in the sample
# offsets inside a pack block: 0 Primary, 1 Carry Fwd, 2 W.S, 3 W.S %, 4 Committed, 5 Total, 6 Total %
FMT = [ACC, ACC, ACC2, PCT, ACC2, ACC, PCT]


def put(ws, rng, value=None, size=18, h="center", wrap=False, fill=None, fmt=None, color=None, border=BOX, bold=True):
    """Write a value into a cell or a merged range (every cell of the range gets the style and the border)."""
    first = rng.split(":")[0]
    got = ws[rng]
    cells = [got] if not isinstance(got, tuple) else [c for row in got for c in (row if isinstance(row, tuple) else (row,))]
    for c in cells:
        c.font = Font(name=config.FONT_NAME, size=size, bold=bold, color=color)
        c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap)
        c.border = border
        if fill:
            c.fill = fill
        if fmt:
            c.number_format = fmt
    ws[first].value = value
    if ":" in rng and rng.split(":")[0] != rng.split(":")[1]:
        ws.merge_cells(rng)


def _fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def _report_sheet(ws, rows, unit, region, year, month):
    for i, w in enumerate(WIDTHS):
        ws.column_dimensions[L(i + 1)].width = w
    n_packs = len(config.WCD_PACKS)
    last_col = L(2 + 7 * n_packs)
    # ---- title and header block ----
    put(ws, f"A1:{last_col}1", config.WCD_TITLE, size=22, fill=TITLE_FILL, color="FFFFFF")
    put(ws, "A2:A4", config.WCD_HEAD_CODE, size=16, wrap=True)
    put(ws, "B2:B4", config.WCD_HEAD_NAME, size=18, wrap=True)
    put(ws, "C2:D2", "Claim Month: ", h="right")
    put(ws, "E2:I2", datetime.datetime(year, month, 1), h="left", size=20, color="FFFF0000", fmt="mmmm\\-yyyy\\.")
    put(ws, "J2:K2", "Unit:", h="right")
    put(ws, "L2:P2", unit, h="left", size=20, color="FFFF0000")
    put(ws, "Q2:R2", "Region:", h="right")
    put(ws, f"S2:{last_col}2", region, h="left", size=20, color="FFFF0000")
    for k, (heading, *_rest) in enumerate(config.WCD_PACKS):
        c0 = 3 + 7 * k
        fill = _fill(BLOCK_FILLS[k % len(BLOCK_FILLS)])
        put(ws, f"{L(c0)}3:{L(c0 + 6)}3", heading, fill=fill)
        for j, text in enumerate(config.WCD_HEADINGS):
            put(ws, f"{L(c0 + j)}4", text, size=13, wrap=True, fill=fill)
    ws.row_dimensions[1].height = ws.row_dimensions[2].height = ws.row_dimensions[3].height = 35.1
    ws.row_dimensions[4].height = 78

    # ---- one row per distributor ----
    first = 5
    last = first + max(len(rows), 1) - 1
    blue = _fill(config.WCD_COLOR_BLUE)
    for i, r in enumerate(rows):
        n = first + i
        ws.row_dimensions[n].height = 90
        put(ws, f"A{n}", int(r.code) if r.code.isdigit() else r.code, size=16)
        put(ws, f"B{n}", r.name, size=16, h="left", wrap=True)
        for k, f in enumerate(r.packs):
            c0 = 3 + 7 * k
            P, CF, W, WP, CM, T, TP = (L(c0 + j) for j in range(7))
            put(ws, f"{P}{n}", f.primary, fmt=ACC)
            put(ws, f"{CF}{n}", 0, fmt=ACC)                                   # Carry Forward Stock: filled by hand
            put(ws, f"{W}{n}", f.ws, fmt=ACC2, h=None)                        # empty stays empty
            put(ws, f"{WP}{n}", f'=IFERROR({W}{n}/(${P}{n}),"**")', fill=YEL, fmt=PCT)
            if f.is_sum_of_two():
                put(ws, f"{CM}{n}", f"={f.card:g}+{f.hth:g}", fmt=ACC2, h=None, fill=blue)
            else:
                put(ws, f"{CM}{n}", f.committed(), fmt=ACC2, h=None)
            put(ws, f"{T}{n}", f"={CM}{n}+{W}{n}", fill=YEL, fmt=ACC)
            put(ws, f"{TP}{n}", f'=IFERROR({T}{n}/(${P}{n}),"**")', fill=YEL, fmt=PCT)
    if not rows:                                                              # keep the sheet valid if nothing was selected
        for j in range(1, 3 + 7 * n_packs):
            put(ws, f"{L(j)}{first}", None)

    # ---- TOTAL row (all distributors) ----
    t = last + 1
    ws.row_dimensions[t].height = 60
    put(ws, f"A{t}:B{t}", "TOTAL", size=20, fill=YEL)
    for k in range(n_packs):
        c0 = 3 + 7 * k
        P, CF, W, WP, CM, T, TP = (L(c0 + j) for j in range(7))
        for col, fmt in ((P, ACC), (CF, ACC), (W, ACC), (CM, ACC), (T, ACC)):
            put(ws, f"{col}{t}", f"=SUM({col}{first}:{col}{last})", fill=YEL, fmt=fmt)
        put(ws, f"{WP}{t}", f'=IFERROR({W}{t}/(${P}{t}),"**")', fill=YEL, fmt=PCT)
        put(ws, f"{TP}{t}", f'=IFERROR({T}{t}/(${P}{t}),"**")', fill=YEL, fmt=PCT)
        # RED: "% age of W.S Claim" above the limit (text such as "**" is ignored)
        red = _fill(config.WCD_COLOR_RED)
        ws.conditional_formatting.add(
            f"{WP}{first}:{WP}{last}",
            FormulaRule(formula=[f"AND(ISNUMBER({WP}{first}),{WP}{first}>{config.WCD_RED_LIMIT})"],
                        fill=PatternFill(start_color=config.WCD_COLOR_RED, end_color=config.WCD_COLOR_RED, fill_type="solid")))
    report_writer._page_setup(ws, t, last_col, "1:4", "landscape",
                              f"{config.WCD_TITLE} - {processing.date_label(year, month)}")


def _check_sheet(ws, issues, run_notes, sources, year, month):
    for col, w in zip("ABCD", (30, 12, 34, 110)):
        ws.column_dimensions[col].width = w
    ws["A1"].value = f"Check sheet - {config.WCD_FILE_WORD} {processing.date_label(year, month)}"
    ws["A1"].font = Font(name=config.FONT_NAME, size=14, bold=True)
    ws["A2"].value = f"Created: {datetime.datetime.now():%d-%b-%Y %H:%M}"
    ws["A2"].font = Font(name=config.FONT_NAME, size=9)
    ws["A3"].value = ("Rules: Primary Lifting = sales file. Variable Disc. W.S = raw cases of the distributor's Wholesale Claim sheet. "
                      "Committed Discount = Card Discount cases + HTH cases (BLUE cell = the sum of two figures, shown as =card+hth). "
                      f"RED % cell = W.S is above {config.WCD_RED_LIMIT:.0%} of Primary Lifting. Carry Forward Stock = 0 (fill by hand). "
                      "Distributors are matched by code. Empty cell = no figure found.")
    ws["A3"].font = Font(name=config.FONT_NAME, size=9)
    thin = Side(style="thin")

    def line(r, vals, fill=None):
        for col, v in zip("ABCD", vals):
            c = ws[f"{col}{r}"]
            c.value = v
            c.font = Font(name=config.FONT_NAME, size=10, bold=fill is None or fill is GREY)
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if fill:
                c.fill = fill

    line(5, ("Distributor", "Code", "Type", "Details"), GREY)
    r = 6
    for kind, text in run_notes:
        line(r, ("(whole run)", "", kind, text), ORANGE if kind.startswith("WARNING") else YELLOW)
        r += 1
    for name, code, kind, text in issues:
        line(r, (name, code, kind, text), ORANGE if kind.startswith("WARNING") else YELLOW)
        r += 1
    if r == 6:
        ws["A6"].value = "No issues found."
        ws["A6"].font = Font(name=config.FONT_NAME, size=10, bold=True)
        r = 7
    r += 1
    ws[f"A{r}"].value = "Files read"
    ws[f"A{r}"].font = Font(name=config.FONT_NAME, size=11, bold=True)
    r += 1
    for s in sources:
        line(r, ("", "", "Source", s))
        r += 1
    report_writer._page_setup(ws, r, "D", "5:5", "landscape", "Check sheet")


def write_wcd_workbook(rows, issues, run_notes, sources, unit, region, year, month, path):
    wb = Workbook()
    ws = wb.active
    ws.title = config.WCD_SHEET_NAME
    _report_sheet(ws, rows, unit, region, year, month)
    _check_sheet(wb.create_sheet("Check"), issues, run_notes, sources, year, month)
    wb.calculation.fullCalcOnLoad = True            # Excel recalculates formulas when opened
    wb.save(path)
