"""
cost_writer.py  -  Draws the "Distribution Wise Cost" workbook (layout copied from your sample):
   sheet "Credit Note" = title, heading row, one row per distributor, Total row, and under it the
                         Sampling / unit block (Sampling Cases ... Per Case Cost)
   sheet "Check"       = everything to verify before using the sheet
Amounts, totals and costs are real Excel formulas (=SUM per row, SUBTOTAL(9, ...) in the Total row, =L/M for the cost).
"""
import calendar
import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from tools.claims import config
from tools.claims import ds_writer
from tools.claims import processing
from tools.claims import report_writer
from tools.claims.ds_writer import ACCOUNTING, GREY, ORANGE, YELLOW, _cell

COST_FMT = "0.0000"
WIDTHS = ds_writer.WIDTHS[:12] + [21.0, 16.1]          # A..L as in the Discount Summary, M and N from the sample (M a little wider for the Arial labels)
COST_FORMULA = '=IFERROR({a}/{b},"****")'              # shows **** instead of an error when Raw Cases is empty / 0


def _summary_sheet(ws, rows, unit, year, month, sampling_cases, sampling_price):
    for i, w in enumerate(WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w
    # ---- title + headings ----
    for i in range(14):
        _cell(ws, f"{chr(65 + i)}1", fill=GREY, size=14, wrap=True)
    # heading = the file name without ".xlsx", e.g. "Muridke Distribution Wise Cost March-2026"
    ws["A1"].value = f"{unit} {config.CO_FILE_WORD} {calendar.month_name[month]}-{year}"
    ws.merge_cells("A1:N1")
    for i, text in enumerate(config.DS_HEADINGS[:12] + config.CO_HEADINGS_EXTRA):
        _cell(ws, f"{chr(65 + i)}2", text, size=10, wrap=True, fill=GREY)
    ws.row_dimensions[1].height = 35.1
    ws.row_dimensions[2].height = 50.1

    # ---- one row per distributor ----
    first = 3
    last = first + len(rows) - 1
    for k, r in enumerate(rows):
        n = first + k
        ws.row_dimensions[n].height = 35.1
        _cell(ws, f"A{n}", k + 1)
        _cell(ws, f"B{n}", int(r.code) if r.code.isdigit() else r.code)
        _cell(ws, f"C{n}", r.name, h="left", wrap=True)
        for col, val in zip("DEFGHIJK", r.amounts):
            _cell(ws, f"{col}{n}", val, h=None, fmt=ACCOUNTING)            # empty stays empty
        _cell(ws, f"L{n}", f"=SUM(D{n}:K{n})", h=None, fmt=ACCOUNTING)
        _cell(ws, f"M{n}", r.raw, h=None, fmt=ACCOUNTING)
        _cell(ws, f"N{n}", COST_FORMULA.format(a=f"L{n}", b=f"M{n}"), fmt=COST_FMT)

    # ---- Total row ----
    t = last + 1
    ws.row_dimensions[t].height = 35.1
    for col in "ABC":
        _cell(ws, f"{col}{t}", fill=GREY)
    ws[f"A{t}"].value = "Total"
    ws.merge_cells(f"A{t}:C{t}")
    for col in "DEFGHIJKLM":
        _cell(ws, f"{col}{t}", f"=SUBTOTAL(9,{col}{first}:{col}{last})", h=None, fill=GREY, fmt=ACCOUNTING)
    _cell(ws, f"N{t}", COST_FORMULA.format(a=f"L{t}", b=f"M{t}"), fill=GREY, fmt=COST_FMT)

    # ---- Sampling / unit block (M:N, two rows under the Total row) ----
    top = t + 2
    labels = config.CO_BLOCK_LABELS
    values = [sampling_cases,                                              # typed by the user
              sampling_price,                                              # typed by the user
              f"=N{top}*N{top + 1}",                                       # Sampling Amount
              f"=L{t}",                                                    # Claims Amount = Total of all distributors
              f"=N{top + 2}+N{top + 3}",                                   # Total Amount
              f"=M{t}",                                                    # Raw Cases of the whole unit
              COST_FORMULA.format(a=f"N{top + 4}", b=f"N{top + 5}")]       # Per Case Cost
    fmts = ["General", "General", ACCOUNTING, ACCOUNTING, ACCOUNTING, ACCOUNTING, COST_FMT]
    for i, (label, val, fmt) in enumerate(zip(labels, values, fmts)):
        r = top + i
        final = i == len(labels) - 1
        ws.row_dimensions[r].height = 20.1
        _cell(ws, f"M{r}", label, h="left", size=14 if final else 12, fill=GREY if final else None)
        _cell(ws, f"N{r}", val, h=None, size=14 if final else 12, fill=GREY if final else None, fmt=fmt)

    end = top + len(labels) - 1
    report_writer._page_setup(ws, end, "N", "1:2", "landscape",
                              f"{unit} {config.CO_FILE_WORD} - {processing.date_label(year, month)}")


def _check_sheet(ws, issues, run_notes, sources, year, month):
    for col, w in zip("ABCD", (30, 12, 34, 110)):
        ws.column_dimensions[col].width = w
    ws["A1"].value = f"Check sheet - {config.CO_FILE_WORD} {processing.date_label(year, month)}"
    ws["A1"].font = Font(name=config.FONT_NAME, size=14, bold=True)
    ws["A2"].value = f"Created: {datetime.datetime.now():%d-%b-%Y %H:%M}"
    ws["A2"].font = Font(name=config.FONT_NAME, size=9)
    ws["A3"].value = ("Rules: claim amounts = the Discount Summary file (including the columns filled by hand there). "
                      f"Raw Cases = '{config.CO_COL_RAW_CASES}' of the sales file, matched by code. Per Unit Cost = Total Amount / Raw Cases. "
                      "Per Case Cost (block) = (Sampling Cases x Sampling Price + Claims Amount) / Raw Cases of all distributors.")
    ws["A3"].font = Font(name=config.FONT_NAME, size=9)
    thin = Side(style="thin")

    def line(r, vals, fill=None):
        for col, v in zip("ABCD", vals):
            c = ws[f"{col}{r}"]
            c.value = v
            c.font = Font(name=config.FONT_NAME, size=10, bold=fill is None or fill is GREY)
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if fill:
                c.fill = fill

    line(5, ("Distributor", "Code", "Type", "Details"), GREY)
    r = 6
    for kind, text in run_notes:
        line(r, ("(whole run)", "", kind, text), ORANGE if kind.startswith("WARNING") else YELLOW)
        r += 1
    for name, code, kind, text in issues:
        line(r, (name, code, kind, text), ORANGE if kind.startswith("WARNING") else YELLOW)
        r += 1
    if r == 6:
        ws["A6"].value = "No issues found."
        ws["A6"].font = Font(name=config.FONT_NAME, size=10, bold=True)
        r = 7
    r += 1
    ws[f"A{r}"].value = "Files read / values used"
    ws[f"A{r}"].font = Font(name=config.FONT_NAME, size=11, bold=True)
    r += 1
    for s in sources:
        line(r, ("", "", "Source", s))
        r += 1
    report_writer._page_setup(ws, r, "D", "5:5", "landscape", "Check sheet")


def write_cost_workbook(rows, issues, run_notes, sources, unit, year, month, sampling_cases, sampling_price, path):
    wb = Workbook()
    ws = wb.active
    ws.title = config.CO_SHEET_NAME
    _summary_sheet(ws, rows, unit, year, month, sampling_cases, sampling_price)
    _check_sheet(wb.create_sheet("Check"), issues, run_notes, sources, year, month)
    wb.calculation.fullCalcOnLoad = True            # Excel recalculates formulas when opened
    wb.save(path)
