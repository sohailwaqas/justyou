"""
ds_reader.py  -  Reads the files for the "Distribution Wise Discount Amount" summary (v0.5):
   * HTH Claims PDFs and Card Discount Claims PDFs  (one file with many distributors, or many files)
   * the Freight Charges workbook and the Wholesale Claim workbook that this tool created earlier
Nothing is changed in the original files.  Every problem is reported as a note; nothing is dropped silently.
"""
import datetime
import os
import re
from dataclasses import dataclass, field

from tools.claims import config
from tools.claims import readers
from tools.claims.readers import FriendlyError

MON3 = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
_CODE = re.compile(r"\((\d{5,7})\)")
_NUM = r"-?[\d,]+(?:\.\d+)?"


def _num(text):
    return float(text.replace(",", ""))


@dataclass
class PdfEntry:
    kind: str            # "HTH" or "CARD"
    code: str
    name: str
    amount: float        # the printed "Total:" amount of the page
    file: str
    page: int
    month: tuple = None  # (month number, year) printed on the page, e.g. (3, 2026)
    lines_sum: float = 0.0
    lines_count: int = 0
    discount_amount: float = None   # card pages only: "Discount Amount <Mon>" (cross-check)


@dataclass
class PdfResult:
    entries: list = field(default_factory=list)     # [PdfEntry]
    grand_total: dict = field(default_factory=dict)  # file name -> printed grand total (card report)
    notes: list = field(default_factory=list)       # [(type, details)]  type starting with "WARNING" is highlighted
    files: list = field(default_factory=list)       # [(file name, text summary)]


def _page_name(lines, kind_idx):
    """Distributor name = the text before '(code)' under the title; it can wrap over two lines."""
    parts = []
    for line in lines[kind_idx + 1:kind_idx + 4]:
        line = line.replace("Copy for IT", " ").strip()
        m = _CODE.search(line)
        if m:
            parts.append(line[:m.start()])
            return m.group(1), " ".join(" ".join(parts).split())
        parts.append(line)
    return None, ""


def read_pdfs(paths):
    """Read every PDF; returns PdfResult. Raises FriendlyError only when nothing usable could be read."""
    try:
        import pdfplumber
    except ImportError:
        raise FriendlyError("The PDF reading library is missing.\nFix: run INSTALL_REQUIREMENTS.bat "
                            "(or: pip install pdfplumber).")
    res = PdfResult()
    for path in paths:
        fname = os.path.basename(path)
        if not os.path.isfile(path) or not path.lower().endswith(".pdf"):
            res.notes.append(("WARNING File skipped", f"'{fname}' is not a PDF file, so it was skipped."))
            res.files.append((fname, "skipped (not a PDF)"))
            continue
        found = 0
        try:
            with pdfplumber.open(path) as pdf:
                for pno, page in enumerate(pdf.pages, 1):
                    text = (page.extract_text() or "").strip()
                    if not text:
                        continue                                  # blank page (the HTH files have one): fine
                    lines = [l.strip() for l in text.split("\n") if l.strip()]
                    title = next((i for i, l in enumerate(lines) if l.startswith("Distributor Credit Note")), None)
                    if title is None:
                        m = re.match(rf"^Grand Total\s+{_NUM}\s+({_NUM})$", text)
                        if m:
                            res.grand_total[fname] = _num(m.group(1))
                        else:
                            res.notes.append(("WARNING Page not recognised",
                                              f"{fname}, page {pno}: this page is not an HTH / Card Discount page "
                                              "and was NOT read. Please check it."))
                        continue
                    if config.DS_PDF_CARD_KEY in lines[title]:
                        kind = "CARD"
                    elif config.DS_PDF_HTH_KEY in lines[title]:
                        kind = "HTH"
                    else:
                        res.notes.append(("WARNING Page not recognised",
                                          f"{fname}, page {pno}: the title '{lines[title]}' is neither HTH nor Card "
                                          "Discount. The page was NOT read."))
                        continue
                    code, name = _page_name(lines, title)
                    total = next((l for l in lines if l.startswith("Total:")), None)
                    tm = re.match(rf"^Total:\s+({_NUM})\s+({_NUM})$", total or "")
                    if not code or not tm:
                        res.notes.append(("WARNING Page not read",
                                          f"{fname}, page {pno} ({kind}): could not find the distributor code or the "
                                          "'Total:' line. The page was NOT read. Please check it."))
                        continue
                    ent = PdfEntry(kind, code, name, _num(tm.group(2)), fname, pno)
                    mm = re.search(r"Month:\s*([A-Za-z]{3})-(\d{4})", text)
                    if mm and mm.group(1).title() in MON3:
                        ent.month = (MON3.index(mm.group(1).title()) + 1, int(mm.group(2)))
                    for l in lines:                               # package lines: "<pack> <cases> <amount>"
                        pm = re.match(rf"^(\d.*?|\d+\.\dL.*?)\s+({_NUM})\s+({_NUM})$", l)
                        if pm and not l.startswith("Total:"):
                            ent.lines_sum += _num(pm.group(3))
                            ent.lines_count += 1
                    dm = re.search(rf"^Discount Amount \w+\s+({_NUM})$", text, re.M)
                    if dm:
                        ent.discount_amount = _num(dm.group(1))
                    res.entries.append(ent)
                    found += 1
        except Exception as e:
            res.notes.append(("WARNING File not read", f"'{fname}' could not be opened as a PDF ({e}). It was skipped."))
            res.files.append((fname, "could not be read"))
            continue
        if found == 0:
            res.notes.append(("WARNING No data found",
                              f"'{fname}' has no readable HTH / Card Discount pages (a scanned picture instead of "
                              "text, or another kind of report?)."))
        kinds = sorted({e.kind for e in res.entries if e.file == fname})
        res.files.append((fname, f"{found} distributor page(s)" + (f" ({'/'.join(kinds)})" if kinds else "")))
    if not res.entries:
        raise FriendlyError("No HTH or Card Discount data could be read from the PDF file(s).\n"
                            + "\n".join(f"- {t}: {d}" for t, d in res.notes))
    return res


# ----------------------------------------------------------------------------- workbooks made by this tool
@dataclass
class SheetAmount:
    code: str
    name: str
    amount: float
    sheet: str
    month: tuple = None      # (month, year) found on the sheet


def _open_book(path, what):
    import openpyxl
    if not path or not os.path.isfile(path):
        raise FriendlyError(f"{what} file not found:\n{path}")
    if not path.lower().endswith(".xlsx"):
        raise FriendlyError(f"The {what} file must be an .xlsx file created by this tool.")
    try:
        return openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True)
    except PermissionError:
        raise FriendlyError(f"The {what} file is locked. Close it in Excel and try again.")
    except Exception as e:
        raise FriendlyError(f"Could not open the {what} file. It may be damaged.\nDetails: {e}")


def _is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def read_freight_book(path):
    """Freight Charges workbook -> ([SheetAmount], notes). Amount = sum of Cases x Rate (computed here, so it
    does not matter whether Excel has already calculated the file)."""
    wb, wv = _open_book(path, "Freight Charges")
    out, notes = [], []
    for ws in wb.worksheets:
        if ws.title.lower() == "check":
            continue
        if str(ws["A5"].value).strip() != "Distributor Code":
            notes.append(("WARNING Sheet not recognised", f"Freight file, sheet '{ws.title}': not a Freight Charges "
                                                          "sheet (heading 'Distributor Code' not found); skipped."))
            continue
        code = str(ws["A6"].value).strip().split(".")[0]
        total = 0.0
        r = 6
        while r <= ws.max_row and str(ws[f"A{r}"].value).strip() != "Total":
            rate, cases = ws[f"D{r}"].value, ws[f"E{r}"].value
            if _is_num(rate) and _is_num(cases):
                total += rate * cases
            r += 1
        cached = wv[ws.title][f"F{r}"].value
        if _is_num(cached) and abs(cached - total) > 0.5:
            notes.append(("WARNING Freight total differs",
                          f"Sheet '{ws.title}': the total in the file ({cached:,.0f}) is not Cases x Rate "
                          f"({total:,.0f}). The file may have been edited by hand. {total:,.0f} is used."))
        d = ws["B4"].value
        month = (d.month, d.year) if isinstance(d, (datetime.date, datetime.datetime)) else None
        out.append(SheetAmount(code, str(ws["E3"].value or "").strip(), total, ws.title, month))
    if not out:
        raise FriendlyError("No Freight Charges sheets were found in the Freight Charges file.")
    return out, notes


def read_wholesale_book(path):
    """Wholesale Claim workbook -> ([SheetAmount], notes). Amount = sum of Raw Cases x Discount Rate of every
    outlet row (the same as the sheet's Net Total)."""
    wb, wv = _open_book(path, "Wholesale Claim")
    out, notes = [], []
    for ws in wb.worksheets:
        if ws.title.lower() == "check":
            continue
        if str(ws["B2"].value).strip() != config.WS_REPORT_TITLE:
            notes.append(("WARNING Sheet not recognised", f"Wholesale file, sheet '{ws.title}': not a Wholesale Claim "
                                                          "sheet (title not found); skipped."))
            continue
        head = str(ws["C4"].value or "")
        m = _CODE.search(head)
        if not m:
            notes.append(("WARNING Sheet not read", f"Wholesale file, sheet '{ws.title}': no '(code) name' in C4; skipped."))
            continue
        sub = net = None
        for r in range(6, ws.max_row + 1):
            v = str(ws[f"B{r}"].value).strip()
            if v == "Sub Total" and sub is None:
                sub = r
            if v == "Net Total":
                net = r
                break
        if sub is None or net is None:
            notes.append(("WARNING Sheet not read", f"Wholesale file, sheet '{ws.title}': 'Sub Total' / 'Net Total' "
                                                    "rows not found; skipped."))
            continue
        total = sum(ws[f"G{r}"].value * ws[f"H{r}"].value for r in range(6, sub)
                    if _is_num(ws[f"G{r}"].value) and _is_num(ws[f"H{r}"].value))
        cached = wv[ws.title][f"I{net}"].value
        if _is_num(cached) and abs(cached - total) > 0.5:
            notes.append(("WARNING Wholesale total differs",
                          f"Sheet '{ws.title}': Net Total in the file ({cached:,.0f}) is not the sum of the outlet "
                          f"rows ({total:,.0f}). The file may have been edited by hand. {total:,.0f} is used."))
        mm = re.search(r"01-([A-Za-z]{3})-(\d{4})", str(ws["F4"].value or ""))
        month = (MON3.index(mm.group(1).title()) + 1, int(mm.group(2))) if mm and mm.group(1).title() in MON3 else None
        name = head[m.end():].strip()
        out.append(SheetAmount(m.group(1), name, total, ws.title, month))
    if not out:
        raise FriendlyError("No Wholesale Claim sheets were found in the Wholesale Claim file.")
    return out, notes
