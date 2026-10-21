"""
ws_writer.py  -  Draws the printable Wholesale Claim sheets (layout copied from Format_Wholesale_Claim)
and the "Check" sheet. Amounts, totals and averages are real Excel formulas.

Sheet layout (one sheet per distributor):
   rows 2-5      title, header block, column headings
   outlet blocks one block per outlet; inside it one row per CONDITION, grouped by pack
   Sub Total     one row per condition (all outlets added up)
   Net Total     everything added up; "Total Amount Rupees" in the header links to it
   signatures    two lines of three signature boxes
"""
import calendar
import datetime
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break

from tools.claims import config
from tools.claims import processing
from tools.claims import report_writer

MED = Side(style="medium")
BOX = Border(left=MED, right=MED, top=MED, bottom=MED)
GREY = PatternFill("solid", fgColor="C0C0C0")
YELLOW = PatternFill("solid", fgColor="FFFF99")
ORANGE = PatternFill("solid", fgColor="FFC000")
NUM = "#,##0"
AVG = "0.0000"
ROW_H = 39.95


def _style(c, size=14, h="center", wrap=False, fill=None, fmt=None, box=True):
    c.font = Font(name=config.FONT_NAME, size=size, bold=True)
    c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap, indent=1 if h == "left" else 0)
    if box:
        c.border = BOX
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt


def _cells(ws, rng):
    """All cells of 'A1' or 'A1:C3' as one flat list."""
    got = ws[rng]
    if not isinstance(got, tuple):
        return [got]
    return [c for row in got for c in (row if isinstance(row, tuple) else (row,))]


def put(ws, rng, value=None, **kw):
    """Write a value into a cell or a merged range (every cell of the range gets the style and border)."""
    first = rng.split(":")[0]
    for c in _cells(ws, rng):
        _style(c, **kw)
    ws[first].value = value
    if ":" in rng and rng.split(":")[0] != rng.split(":")[1]:
        ws.merge_cells(rng)


def _lines(text, per_line):
    """Rough number of lines a text needs when wrapped at 'per_line' characters (words are not split)."""
    total = 0
    for part in str(text).split("\n"):
        n, cur = 1, 0
        for word in part.split():
            w = len(word)
            while w > per_line:                       # a very long word is broken
                n += 1
                w -= per_line
            if cur and cur + 1 + w > per_line:
                n, cur = n + 1, w
            else:
                cur = cur + w + (1 if cur else 0)
        total += n
    return total


def _outlet_text(o):
    return f"{o.name}\n{o.address}" if (config.WS_SHOW_ADDRESS and o.address) else o.name


def _distributor_sheet(ws, rep, conds, year, month):
    C = len(conds)
    n = len(rep.outlets)
    for col, w in zip("ABCDEFGHIJKLM", (1.7, 15.56, 42.56, 15.7, 13.56, 14.56, 14.56, 14.56, 14.56, 14.56, 14.56, 14.28, 10.85)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["M"].hidden = True                 # hidden check column (TRUE = block adds up)

    # pack blocks inside one outlet: [(first offset, last offset, label)]
    blocks, i = [], 0
    while i < C:
        j = i
        while j + 1 < C and conds[j + 1].pack.upper() == conds[i].pack.upper():
            j += 1
        blocks.append((i, j, conds[i].pack_label()))
        i = j + 1

    # row heights: grow when the brand text needs more lines
    cond_h = [max(ROW_H, 19 * _lines(c.brand_text(), 8) + 6) for c in conds]
    sb = 6 + n * C                 # first Sub Total row
    nt = sb + C                    # Net Total row
    ws.row_dimensions[1].height = 5.1
    for r in range(2, 6):
        ws.row_dimensions[r].height = ROW_H
    for o, out in enumerate(rep.outlets):
        need = 19 * _lines(_outlet_text(out), 24) + 8          # the outlet cell must show all its text
        extra = max(0.0, (need - sum(cond_h)) / C)
        for k in range(C):
            ws.row_dimensions[6 + o * C + k].height = cond_h[k] + extra
    ws.row_dimensions[4].height = max(ROW_H, 21 * _lines(f"({rep.code}) {rep.name}", 24) + 6)
    ws.row_dimensions[3].height = max(ROW_H, 21 * _lines(rep.unit, 24) + 6)
    for k in range(C):
        ws.row_dimensions[sb + k].height = cond_h[k]
    ws.row_dimensions[nt].height = ROW_H

    # ---- title + header block ----
    last_day = calendar.monthrange(year, month)[1]
    mon = processing.MONTHS[month - 1]
    put(ws, "B2:L2", config.WS_REPORT_TITLE, size=28)
    put(ws, "B3", "Unit:", size=16, h="right")
    put(ws, "C3", rep.unit, size=16, h="left", wrap=True)
    put(ws, "D3:E3", "SDM/O/S:", size=16, h="right")
    put(ws, "F3:I3", rep.sdm, size=16, h="left", wrap=True)
    put(ws, "J3:J4", "Total Amount Rupees:", size=16, wrap=True)
    put(ws, "K3:L4", f"=I{nt}", size=20, fmt=NUM)
    put(ws, "B4", "Distributor:\n& Code", size=12, h="right", wrap=True)
    put(ws, "C4", f"({rep.code}) {rep.name}", size=16, h="left", wrap=True)
    put(ws, "D4:E4", "Month/Year:", size=16, h="right")
    put(ws, "F4:I4", f"01-{mon}-{year} to {last_day:02d}-{mon}-{year}", size=16)
    for col, text, size in (("B", "Code", 16), ("C", "Outlet Name & Address", 16), ("D", "Pack", 16), ("E", "Brand", 16),
                            ("F", "Claim\nDate", 14), ("G", "Discount\nRate", 14), ("H", "Raw\nCases", 14),
                            ("I", "Amount", 14), ("J", "Total\nCases", 14), ("K", "Total\nAmount", 14),
                            ("L", "Average\nRate", 14)):
        put(ws, f"{col}5", text, size=size, wrap=True)

    def pack_totals(top, fill=None):
        """Total Cases / Total Amount / Average Rate for each pack block that starts at row 'top'."""
        for a, b, _ in blocks:
            r1, r2 = top + a, top + b
            rng = lambda col: f"{col}{r1}:{col}{r2}" if r2 > r1 else f"{col}{r1}"
            put(ws, rng("J"), f"=SUM(H{r1}:H{r2})", size=16, fmt=NUM, fill=fill)
            put(ws, rng("K"), f"=SUM(I{r1}:I{r2})", size=16, fmt=NUM, fill=fill)
            put(ws, rng("L"), f'=IFERROR(K{r1}/J{r1},"****")', size=16, fmt=AVG, fill=fill)
            put(ws, f"M{r1}", f"=SUM(H{r1}:H{r2})=J{r1}", size=14)

    def pack_cells(top, fill=None):
        for a, b, label in blocks:
            r1, r2 = top + a, top + b
            put(ws, f"D{r1}:D{r2}" if r2 > r1 else f"D{r1}", label, size=14, wrap=True, fill=fill)

    # ---- outlet blocks ----
    for o, out in enumerate(rep.outlets):
        top = 6 + o * C
        bot = top + C - 1
        put(ws, f"B{top}:B{bot}" if C > 1 else f"B{top}", int(out.code) if out.code.isdigit() else out.code, size=11, wrap=True)
        put(ws, f"C{top}:C{bot}" if C > 1 else f"C{top}", _outlet_text(out),
            size=14 if (config.WS_SHOW_ADDRESS and out.address) else 18, wrap=True)
        pack_cells(top)
        for k, (cond, cs) in enumerate(zip(conds, out.cases)):
            r = top + k
            put(ws, f"E{r}", cond.brand_text(), size=14, wrap=True)
            put(ws, f"F{r}", cond.date_text(month), size=14, wrap=True)
            put(ws, f"G{r}", cond.rate, size=16)
            put(ws, f"H{r}", cs, size=16, fmt=NUM)
            put(ws, f"I{r}", f"=H{r}*G{r}", size=16, fmt=NUM)
        pack_totals(top)

    # ---- Sub Total (one row per condition, all outlets added up) ----
    put(ws, f"B{sb}:C{nt - 1}", "Sub Total", size=20, fill=GREY)        # B and C merged so the text is fully visible
    pack_cells(sb, GREY)
    for k, cond in enumerate(conds):
        r = sb + k
        refs_h = ",".join(f"H{6 + o * C + k}" for o in range(n))
        refs_i = ",".join(f"I{6 + o * C + k}" for o in range(n))
        put(ws, f"E{r}", cond.brand_text(), size=14, wrap=True, fill=GREY)
        put(ws, f"F{r}", cond.date_text(month), size=14, wrap=True, fill=GREY)
        put(ws, f"G{r}", f"=IFERROR(I{r}/H{r},{cond.rate:g})", size=16, fill=GREY, fmt="General")
        put(ws, f"H{r}", f"=SUM({refs_h})" if n else 0, size=16, fmt=NUM, fill=GREY)
        put(ws, f"I{r}", f"=SUM({refs_i})" if n else 0, size=16, fmt=NUM, fill=GREY)
    pack_totals(sb, GREY)

    # ---- Net Total ----
    put(ws, f"B{nt}:G{nt}", "Net Total", size=18)
    put(ws, f"H{nt}", f"=SUM(H{sb}:H{nt - 1})", size=16, fmt=NUM)
    put(ws, f"I{nt}", f"=SUM(I{sb}:I{nt - 1})", size=16, fmt=NUM)
    put(ws, f"J{nt}", f"=H{nt}", size=16, fmt=NUM)
    put(ws, f"K{nt}", f"=I{nt}", size=16, fmt=NUM)
    put(ws, f"L{nt}", f'=IFERROR(K{nt}/J{nt},"****")', size=16, fmt=AVG)

    # ---- signatures ----
    def signatures(line_row, labels):
        for rr in range(line_row - 1, line_row + 2):
            ws.row_dimensions[rr].height = 30
        for rng, text in zip(("C{r}", "E{r}:H{r}", "J{r}:L{r}"), labels):
            line, lab = rng.format(r=line_row), rng.format(r=line_row + 1)
            for c in _cells(ws, line):
                c.border = Border(bottom=MED)
            put(ws, lab, text, size=16, box=False)
            for c in _cells(ws, lab):
                c.border = Border(top=MED) if rng != "C{r}" else Border()

    signatures(nt + 4, config.WS_SIGNATURES_ROW_1)
    signatures(nt + 9, config.WS_SIGNATURES_ROW_2)
    last_row = nt + 10
    report_writer._page_setup(ws, last_row, "L", "1:5", config.ORIENTATION,
                              f"{config.WS_REPORT_TITLE} - {processing.date_label(year, month)}")
    ws.print_area = f"B2:L{last_row}"
    _page_breaks(ws, [(6 + o * C, 6 + o * C + C - 1) for o in range(n)], (sb, last_row))


def _page_breaks(ws, outlet_blocks, tail):
    """Add page breaks so that no outlet block is cut in half (A4 portrait, fitted to one page wide).
    outlet_blocks = [(first row, last row)], tail = (first row of Sub Total, last row of the sheet)."""
    def height(a, b):
        return sum(ws.row_dimensions[r].height or 15 for r in range(a, b + 1))
    width_pt = sum(ws.column_dimensions[c].width for c in "BCDEFGHIJKL") * 5.6        # approx. points
    page_w = 595 - 2 * config.MARGIN_LEFT_RIGHT * 72
    page_h = 842 - 2 * config.MARGIN_TOP_BOTTOM * 72
    scale = min(1.0, page_w / width_pt) / 0.85                                      # assume 15% less room (Excel and LibreOffice differ a little)
    room = page_h / scale - height(1, 5)                                               # body room on every page
    used = height(1, 5) * 0                                                            # page 1 also shows rows 1-5
    for first, last in outlet_blocks + [tail]:
        h = height(first, last) * (1.1 if (first, last) == tail else 1.0)
        if used and used + h > room:
            ws.row_breaks.append(Break(id=first - 1))
            used = 0
        used += h


def _check_sheet(ws, reports, conds, run_notes, year, month, source_name):
    """One line per thing the user should look at before printing. WARNING lines are orange."""
    for col, w in zip("ABCD", (30, 12, 28, 110)):
        ws.column_dimensions[col].width = w
    ws["A1"].value = f"Check sheet - Wholesale Claim {processing.date_label(year, month)}"
    ws["A1"].font = Font(name=config.FONT_NAME, size=14, bold=True)
    ws["A2"].value = f"Source file: {source_name}      Created: {datetime.datetime.now():%d-%b-%Y %H:%M}"
    ws["A2"].font = Font(name=config.FONT_NAME, size=9)
    ws["A3"].value = "Rule: a sale is counted by its BOOKING date (ORDER_DATE). Delivery date is only used for warnings."
    ws["A3"].font = Font(name=config.FONT_NAME, size=9)
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    def line(r, vals, fill=None):
        for col, v in zip("ABCD", vals):
            c = ws[f"{col}{r}"]
            c.value = v
            c.font = Font(name=config.FONT_NAME, size=10, bold=fill is None)
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c.border = box
            if fill:
                c.fill = fill

    line(5, ("Distributor", "Code", "Type", "Details"), GREY)
    r = 6
    for text in run_notes:
        kind = "WARNING Overlapping conditions" if text.startswith("WARNING Overlapping") else "File note"
        text = text.split(": ", 1)[1] if kind != "File note" else text
        line(r, ("(whole run)", "", kind, text), ORANGE if kind.startswith("WARNING") else YELLOW)
        r += 1
    for rep in reports:
        for kind, text in rep.issues:
            line(r, (rep.name, rep.code, kind, text), ORANGE if kind.startswith("WARNING") else YELLOW)
            r += 1
    if r == 6:
        ws["A6"].value = "No issues found."
        ws["A6"].font = Font(name=config.FONT_NAME, size=10, bold=True)
        r = 7
    r += 1
    ws[f"A{r}"].value = "Conditions used"
    ws[f"A{r}"].font = Font(name=config.FONT_NAME, size=11, bold=True)
    r += 1
    for k, c in enumerate(conds, 1):
        line(r, (f"Condition {k}", "", f"Rs {c.rate:g} per case",
                 f"Pack: {c.pack_label()}   |   Brand: {c.brand_text()}   |   Claim date: {c.date_text(month)}"), None)
        r += 1
    report_writer._page_setup(ws, r, "D", "5:5", "landscape", "Check sheet")


def write_ws_workbook(reports, conds, year, month, path, source_name="", run_notes=()):
    wb = Workbook()
    wb.remove(wb.active)
    for rep in reports:
        _distributor_sheet(wb.create_sheet(rep.sheet_name), rep, conds, year, month)
    _check_sheet(wb.create_sheet("Check"), reports, conds, list(run_notes), year, month, source_name)
    wb.calculation.fullCalcOnLoad = True            # Excel recalculates formulas when opened
    wb.save(path)
