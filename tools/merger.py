"""Merger - logic module (no Streamlit import).

Core logic is the original merge_distributions_gui.py, with two edits:
  1. merge() accepts an optional `summary_out` list that receives the per-file
     summary tuples (file, dist code, rows in sheet, rows kept).
  2. The row-by-row ws.delete_rows() loop is replaced by delete_rows_except(),
     a one-pass equivalent (the original is quadratic and hangs on large files).
The tkinter GUI, threading, queue and 'open folder' code are removed.

Added on top of the original logic:
  merge_bytes()  - in-memory wrapper for the web app (bytes in, bytes out)
  CLI            - python -m tools.merger --unit U --nature N --month M --year Y -o DIR files...

For every file (in the order given):
  1. The chosen sheet (default: first sheet) is used; filters are removed and
     hidden columns / rows are unhidden
  2. Links are broken: formulas (incl. links to other workbooks) become plain
     values, hyperlinks and external-link tables are removed - data is kept
  3. The Dist. Code (taken from the file name) is selected in the 'Dist. Code
     Header' column and every row with any other code is deleted
  4. The remaining rows are appended to the first file's data
  5. All rows get identical formatting (style of the first file's first data
     row, uniform row height, widest column widths across files)
"""
from __future__ import annotations

import copy
import re
import tempfile
from pathlib import Path
from typing import Callable, Optional

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]


class MergeError(Exception):
    pass


# ============================================================ core logic
def norm(v):
    return re.sub(r"[^a-z0-9]+", " ", str(v or "").lower()).strip()


def clean_filename_part(s):
    return re.sub(r'[\\/:*?"<>|]+', "-", s).strip()


def remove_filters(ws):
    ws.auto_filter.ref = None
    ws.auto_filter.filterColumn = []
    ws.auto_filter.sortState = None
    for tbl in ws.tables.values():
        tbl.autoFilter = None
    for dim in ws.row_dimensions.values():
        dim.hidden = False


def unhide_all(ws):
    """Unhide every hidden column (incl. grouped/collapsed) and every hidden row.
    Returns a list of the column letters that were hidden."""
    unhidden = []
    for key, dim in ws.column_dimensions.items():
        if dim.hidden:
            lo = dim.min or 1
            hi = dim.max or lo
            unhidden.extend(get_column_letter(c) for c in range(lo, hi + 1))
            dim.hidden = False
    for dim in ws.row_dimensions.values():
        dim.hidden = False
    return unhidden


def collect_widths(ws, widths):
    """Record the widest width seen for every column index across all files."""
    for dim in ws.column_dimensions.values():
        if dim.width is None:
            continue
        lo = dim.min or 1
        hi = dim.max or lo
        for c in range(lo, hi + 1):
            widths[c] = max(widths.get(c, 0), dim.width if dim.width > 1 else 12)


def normalise_format(ws, hdr_row, last_row, max_col, widths):
    """Give every data row (all distributions) exactly the same formatting:
    cell styles come from the first data row of the first file, row height is
    uniform, and column widths are the widest seen in any file."""
    first = hdr_row + 1
    if last_row >= first:
        template = [copy.copy(ws.cell(first, c)._style) for c in range(1, max_col + 1)]
        height = ws.row_dimensions[first].height
        for r in range(first, last_row + 1):
            for c in range(1, max_col + 1):
                ws.cell(r, c)._style = copy.copy(template[c - 1])
            ws.row_dimensions[r].height = height
    ws.column_dimensions.clear()
    for c, w in widths.items():
        d = ws.column_dimensions[get_column_letter(c)]
        d.width = w
        d.hidden = False


CODE_HEADER_FALLBACKS = ("dist code", "distribution code", "dist", "distribution")


def find_header(ws, code_header, header_row=None, scan=20):
    """Locate the Dist. Code column. Tries the header the user typed first
    (exact cell match, then 'contains'), then common alternatives.
    Returns (header_row, column_index) - both 1-based."""
    keys = [norm(code_header)] + [k for k in CODE_HEADER_FALLBACKS if k != norm(code_header)]
    rows = [header_row] if header_row else range(1, min(ws.max_row, scan) + 1)
    for key in keys:
        if not key:
            continue
        for exact in (True, False):
            for r in rows:
                for c in range(1, ws.max_column + 1):
                    t = norm(ws.cell(r, c).value)
                    if t and (t == key if exact else key in t):
                        return r, c
    return None, None


def dist_from_filename(path):
    """'LESCO.xlsx' / 'LESCO - Sept.xlsx' / 'LESCO_2026.xlsx' -> 'LESCO'"""
    return re.split(r"\s+-\s+|_", Path(path).stem)[0].strip()


def norm_code(v):
    """Normalise a distribution code for comparison: 11111, 11111.0 and '11111'
    are equal, and so are '00123' and 123."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = norm(v)
    if t.isdigit():
        t = t.lstrip("0") or "0"
    return t


def row_matches(cell_value, code, loose=False):
    a, b = norm_code(cell_value), norm_code(code)
    if not a or not b:
        return False
    return (b in a or a in b) if loose else a == b


def resolve_code(ws, hdr_row, code_col, path):
    """Work out which code in the Dist. Code column the file name refers to.
    Tries the whole file name, the part before ' - ' / '_', then each word of
    the name (longest first) and returns the first one that really exists in
    the column. Falls back to the part before ' - ' / '_'."""
    present = {norm_code(ws.cell(r, code_col).value)
               for r in range(hdr_row + 1, ws.max_row + 1)}
    stem = Path(path).stem.strip()
    words = sorted(filter(None, re.split(r"[\s_\-]+", stem)), key=len, reverse=True)
    for cand in [stem, dist_from_filename(path), *words]:
        if norm_code(cand) in present:
            return cand
    return dist_from_filename(path)


def kept_rows(ws, hdr_row, dist_col, dist):
    """Rows whose Dist. Code equals the code taken from the file name
    (ignoring case/punctuation). If none match exactly, falls back to
    'contains' matching."""
    data = range(hdr_row + 1, ws.max_row + 1)
    rows = [r for r in data if row_matches(ws.cell(r, dist_col).value, dist)]
    if not rows:
        rows = [r for r in data if row_matches(ws.cell(r, dist_col).value, dist, True)]
    return rows


# ---- breaking links ------------------------------------------------------
EXTERNAL_REF = re.compile(r"\[\d+\]|\[[^\]]*\.xl[a-z]*\]", re.I)


def break_links(wb, ws, ws_values):
    """Replace every formula in `ws` with its last calculated value, remove
    hyperlinks, and drop the workbook's external-link table, so the data stays
    but nothing points to another file/cell any more.
    `ws_values` is the same sheet loaded with data_only=True (cached values).
    Returns (external_formulas, other_formulas, hyperlinks, formulas_without_value)."""
    ext = other = links = missing = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                v = cell.value
                text = v if isinstance(v, str) else (getattr(v, "text", "") or "")
                cached = ws_values[cell.coordinate].value
                if EXTERNAL_REF.search(text):
                    ext += 1
                else:
                    other += 1
                if cached is None:
                    missing += 1
                cell.value = cached
            if cell.hyperlink is not None:
                cell.hyperlink = None
                links += 1
    ws._hyperlinks = []
    wb._external_links = []
    try:                                  # defined names pointing to other files
        for name in list(wb.defined_names):
            if "[" in (wb.defined_names[name].attr_text or ""):
                del wb.defined_names[name]
    except Exception:
        pass
    return ext, other, links, missing


def copy_row(src_ws, src_r, dst_ws, dst_r, max_col):
    for c in range(1, max_col + 1):
        s, d = src_ws.cell(src_r, c), dst_ws.cell(dst_r, c)
        d.value = s.value
        if s.has_style:
            d._style = copy.copy(s._style)
    h = src_ws.row_dimensions[src_r].height
    if h:
        dst_ws.row_dimensions[dst_r].height = h


def pick_sheet(wb, sheet_name, fname):
    """Blank sheet_name -> first sheet; otherwise the sheet with that name."""
    if not sheet_name:
        return wb.worksheets[0]
    for ws in wb.worksheets:
        if ws.title.strip().lower() == sheet_name.strip().lower():
            return ws
    raise MergeError(f"'{fname}' has no sheet named '{sheet_name}'.\n"
                     f"Sheets in this file: {', '.join(wb.sheetnames)}\n"
                     f"Fix the sheet name, or leave it blank to use the first sheet.")


def delete_rows_except(ws, hdr_row, keep):
    """Remove every data row below `hdr_row` that is not in `keep`, in ONE pass.
    Same result as calling ws.delete_rows(r) bottom-up for each unwanted row
    (cells shift up, row heights/merges are not moved), but linear instead of
    quadratic - openpyxl's delete_rows shifts all cells below on every call,
    which never finishes on large sheets."""
    new_index = {r: r for r in range(1, hdr_row + 1)}
    for n, r in enumerate(sorted(keep), hdr_row + 1):
        new_index[r] = n
    moved = {}
    for (r, c), cell in ws._cells.items():
        nr = new_index.get(r)
        if nr is not None:
            cell.row = nr
            moved[(nr, c)] = cell
    ws._cells = moved


def merge(files, unit, nature, month, year, out_dir, sheet_name=None,
          dist_col="Dist. Code", header_row=None, log=print, progress=lambda i, n: None,
          summary_out=None):
    base_wb = base_ws = None
    base_hdr = base_col = next_row = None
    summary = []
    widths = {}

    for i, f in enumerate(files, 1):
        f = Path(f)
        log(f"Processing {f.name}")
        try:
            wb = load_workbook(f)
            wb_values = load_workbook(f, data_only=True)
        except Exception as e:
            raise MergeError(f"Cannot open '{f.name}': {e}")
        ws = pick_sheet(wb, sheet_name, f.name)
        log(f"  using sheet: {ws.title}")
        remove_filters(ws)
        hidden_cols = unhide_all(ws)
        if hidden_cols:
            log(f"  unhid column(s): {', '.join(hidden_cols)}")
        collect_widths(ws, widths)

        ext, other, links, missing = break_links(wb, ws, wb_values[ws.title])
        if ext or other or links:
            log(f"  broke links: {ext} external-link formula(s), {other} other formula(s) "
                f"-> values; {links} hyperlink(s) removed")
        if missing:
            log(f"  WARNING: {missing} formula(s) had no saved value (file never "
                f"calculated?) - those cells are now empty. Open/save the file in "
                f"Excel once and retry.")

        hdr, col = find_header(ws, dist_col, header_row)
        if not hdr:
            raise MergeError(f"'{f.name}': the Dist. Code column ('{dist_col}') was not found.\n"
                             f"Check 'Dist. Code Header' / header row under Advanced options.")

        dist = resolve_code(ws, hdr, col, f)
        log(f"  Dist. Code from file name: {dist}")

        rows = kept_rows(ws, hdr, col, dist)
        total = ws.max_row - hdr
        if not rows:
            log(f"  WARNING: no rows for '{dist}' - file skipped.")
            summary.append((f.name, dist, total, 0))
            progress(i, len(files))
            continue

        if base_wb is None:
            delete_rows_except(ws, hdr, set(rows))    # keep only the matching rows
            base_wb, base_ws, base_hdr, base_col = wb, ws, hdr, col
            next_row = base_ws.max_row + 1
            for name in base_wb.sheetnames:           # keep a single sheet only
                if name != base_ws.title:
                    del base_wb[name]
        else:
            if (hdr, col) != (base_hdr, base_col):
                log(f"  WARNING: header layout differs from first file "
                    f"(row {hdr}/col {col} vs {base_hdr}/{base_col}) - please check.")
            max_col = max(ws.max_column, base_ws.max_column)
            for r in rows:
                copy_row(ws, r, base_ws, next_row, max_col)
                next_row += 1

        summary.append((f.name, dist, total, len(rows)))
        log(f"  kept {len(rows)} of {total} rows")
        progress(i, len(files))

    if base_wb is None:
        raise MergeError("Nothing to merge - no rows matched in any file.")

    remove_filters(base_ws)
    unhide_all(base_ws)
    normalise_format(base_ws, base_hdr, next_row - 1,
                     max(base_ws.max_column, max(widths, default=0)), widths)
    log("Applied uniform formatting to all rows.")
    out_name = clean_filename_part(f"{unit} {nature} {month} {year}")
    base_ws.title = out_name[:31]                     # Excel tab-name limit
    out_path = Path(out_dir) / f"{out_name}.xlsx"
    try:
        base_wb.save(out_path)
    except PermissionError:
        raise MergeError(f"Cannot save '{out_path.name}'. Is it open in Excel? Close it and retry.")

    log(f"\nTotal data rows: {next_row - base_hdr - 1}")
    log(f"Saved: {out_path}")
    if summary_out is not None:                       # web-app addition
        summary_out.extend(summary)
    return out_path



# ========================================================== web wrappers
def merge_bytes(files: list[tuple[str, bytes]], unit: str, nature: str, month: str, year: str,
                sheet_name: Optional[str] = None, dist_col: str = "Dist. Code",
                header_row: Optional[int] = None,
                log: Callable[[str], None] = print,
                progress: Callable[[int, int], None] = lambda i, n: None
                ) -> tuple[str, bytes, list[tuple[str, str, int, int]]]:
    """Run merge() fully in memory.

    `files` is a list of (file name, file bytes) in merge order; the file name
    matters because the Dist. Code is taken from it. Returns
    (output file name, output bytes, summary) where summary rows are
    (file name, dist code, rows in sheet, rows kept). Nothing is stored on disk
    beyond a temporary folder that is deleted before returning.
    """
    with tempfile.TemporaryDirectory(prefix="handy_merger_") as tmp:
        root = Path(tmp)
        paths = []
        for i, (name, data) in enumerate(files):
            folder = root / "in" / str(i)          # own folder: same file names cannot clash
            folder.mkdir(parents=True)
            p = folder / Path(name).name
            p.write_bytes(data)
            paths.append(p)
        out_dir = root / "out"
        out_dir.mkdir()
        summary: list = []
        out_path = merge(paths, unit, nature, month, year, out_dir, sheet_name=sheet_name,
                         dist_col=dist_col, header_row=header_row, log=log,
                         progress=progress, summary_out=summary)
        return out_path.name, out_path.read_bytes(), summary


def main(argv: Optional[list[str]] = None) -> None:
    """Command-line entry point (testing / scripting)."""
    import argparse
    import datetime
    t = datetime.date.today()
    ap = argparse.ArgumentParser(description="Merge per-distribution Excel files.")
    ap.add_argument("files", nargs="+")
    ap.add_argument("--unit", required=True)
    ap.add_argument("--nature", required=True)
    ap.add_argument("--month", default=MONTHS[t.month - 1])
    ap.add_argument("--year", default=str(t.year))
    ap.add_argument("-o", "--out-dir", default=".")
    ap.add_argument("--sheet", default=None, help="blank = first sheet")
    ap.add_argument("--dist-col", default="Dist. Code")
    ap.add_argument("--header-row", type=int, default=None)
    a = ap.parse_args(argv)
    merge(a.files, a.unit, a.nature, a.month, a.year, a.out_dir, sheet_name=a.sheet,
          dist_col=a.dist_col, header_row=a.header_row)


if __name__ == "__main__":
    main()
